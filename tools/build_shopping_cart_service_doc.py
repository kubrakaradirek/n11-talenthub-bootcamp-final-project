from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\Pierm\IdeaProjects\n11-talenthub-bootcamp-final-project\Shopping_Cart_Service_Mulakat_Notlari.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_table(table):
    table.style = "Table Grid"
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)
            if row_idx == 0:
                set_cell_shading(cell, "EAF2F8")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 78, 121)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_qa(doc, question, answer):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 3"]
    p.add_run(question)
    a = doc.add_paragraph()
    a.add_run(answer)


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 22, RGBColor(31, 78, 121)),
        ("Heading 1", 15, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(54, 96, 146)),
        ("Heading 3", 11, RGBColor(0, 0, 0)),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = section.header.paragraphs[0]
    header.text = "Shopping Cart Service - Mülakat Notları"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 120, 120)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Shopping Cart Service")
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 120, 120)


def build():
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Shopping Cart Service")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Servis çalışma mantığı, class açıklamaları, kullanılan teknolojiler ve mülakat soru-cevapları")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(90, 90, 90)

    add_table(
        doc,
        ["Alan", "Bilgi"],
        [
            ["Modül adı", "shopping-card-service"],
            ["Servis adı", "shopping-cart-service"],
            ["Port", "8765"],
            ["Ana endpoint", "/api/shopping-cart"],
            ["Ana görevi", "Kullanıcı bazlı sepeti Redis üzerinde tutmak ve sepet işlemlerini yönetmek"],
            ["Bağımlı servis", "product-service üzerinden ürün fiyatı, başlık, görsel ve renk bilgisi alınır"],
            ["Mesajlaşma", "RabbitMQ shopping_cart_queue kuyruğuna audit mesajı gönderilir ve aynı servis içinde tüketilir"],
        ],
    )

    doc.add_heading("1. Servisin Kısa Özeti", level=1)
    doc.add_paragraph(
        "Shopping Cart Service, e-ticaret projesinde kullanıcıların sepetini yöneten mikroservistir. "
        "Kullanıcının sepeti username anahtarıyla Redis'te saklanır. Ürün sepete eklenirken ürün bilgisi "
        "Product Service'ten alınır, sepet toplamı yeniden hesaplanır ve işlem RabbitMQ kuyruğuna audit amaçlı yazılır."
    )
    add_bullets(
        doc,
        [
            "Kullanıcı sepetini getirir; sepet yoksa boş sepet modeli oluşturur.",
            "Ürün eklerken product-service'e RestTemplate ile istek atar.",
            "Aynı ürün tekrar eklenirse yeni satır açmak yerine quantity değerini artırır.",
            "Ürün adedi güncellenebilir, ürün sepetten çıkarılabilir ve sepet tamamen temizlenebilir.",
            "Her ekleme sonrası RabbitMQ'ya basit bir audit mesajı publish edilir.",
        ],
    )

    doc.add_heading("2. Servis Nasıl Çalışır?", level=1)
    doc.add_heading("2.1 Genel Akış", level=2)
    add_numbered(
        doc,
        [
            "Frontend veya API Gateway /api/shopping-cart/{username}/add endpoint'ine istek gönderir.",
            "Controller isteği ShoppingCardService katmanına devreder.",
            "Service önce Redis'ten kullanıcının mevcut sepetini alır; yoksa boş sepet oluşturur.",
            "Ürün fiyatı ve bilgileri product-service üzerinden okunur.",
            "Sepette ürün varsa miktar artırılır, yoksa CardItem olarak listeye eklenir.",
            "Toplam fiyat price * quantity üzerinden tekrar hesaplanır.",
            "Güncel sepet Redis'e kaydedilir.",
            "RabbitMQ shopping_cart_queue kuyruğuna audit mesajı gönderilir.",
            "CartAuditConsumer mesajı tüketip log'a yazar.",
        ],
    )

    doc.add_heading("2.2 Endpointler", level=2)
    add_table(
        doc,
        ["Method", "Endpoint", "Ne yapar?", "Dönen sonuç"],
        [
            ["GET", "/api/shopping-cart/{username}", "Kullanıcının sepetini getirir.", "ShoppingCard JSON"],
            ["POST", "/api/shopping-cart/{username}/add?productId=&quantity=", "Ürünü sepete ekler veya mevcut ürünün miktarını artırır.", "Güncel sepet"],
            ["POST", "/api/shopping-cart/{username}/update?productId=&quantity=", "Sepetteki ürünün miktarını verilen quantity değerine çeker.", "Güncel sepet"],
            ["DELETE", "/api/shopping-cart/{username}/remove?productId=", "Ürünü sepet listesinden çıkarır.", "Güncel sepet"],
            ["DELETE", "/api/shopping-cart/{username}/clear", "Kullanıcının sepetini Redis'ten siler.", "204 No Content"],
        ],
    )

    doc.add_heading("2.3 Order Service ile İlişkisi", level=2)
    doc.add_paragraph(
        "Kullanıcı siparişi tamamladığında ödeme ve sipariş akışından sonra sepet temizlenir. "
        "Bu yüzden shopping-cart-service, order-service'ten önceki alışveriş niyeti/verisi aşamasıdır. "
        "Order Service kalıcı sipariş kaydı ve stok/ödeme süreçleriyle ilgilenirken, Shopping Cart Service geçici sepet durumunu tutar."
    )

    doc.add_heading("3. Class Bazlı Açıklama", level=1)
    add_table(
        doc,
        ["Class / Dosya", "Katman", "Görevi", "Kullandığı teknoloji"],
        [
            ["ShoppingCardServiceApplication", "Bootstrapping", "Spring Boot uygulamasını başlatır, Eureka discovery client'ı aktif eder ve load-balanced RestTemplate bean'i üretir.", "Spring Boot, Eureka, RestTemplate, LoadBalancer"],
            ["ShoppingCardController", "Controller", "HTTP endpointlerini tanımlar; get, add, update, remove ve clear işlemlerini service katmanına yönlendirir.", "Spring Web MVC, Swagger/OpenAPI, Lombok"],
            ["ShoppingCardService", "Business Service", "Sepet iş kurallarını uygular; Redis'ten sepet okur/yazar, product-service'ten ürün bilgisi alır, toplam fiyatı hesaplar, RabbitMQ'ya mesaj gönderir.", "Spring Service, Redis Repository, RestTemplate, RabbitTemplate, SLF4J"],
            ["ShoppingCard", "Model / Redis Entity", "Bir kullanıcının sepetini temsil eder. username id olarak kullanılır, items listesi ve totalPrice taşır.", "Spring Data Redis, Serializable, Lombok"],
            ["CardItem", "Model", "Sepetteki tek ürün satırını temsil eder: productId, title, quantity, price, imageUrl, color.", "Serializable, Lombok"],
            ["ShoppingCardRepository", "Repository", "ShoppingCard nesneleri için Redis üzerinde CRUD operasyonları sağlar.", "Spring Data CrudRepository, Redis"],
            ["ProductRepository", "Repository", "Dosya mevcut ama içinde repository tanımı yok; aktif olarak kullanılmıyor.", "Yok / temizlenebilir"],
            ["RabbitMQConfig", "Configuration", "shopping_cart_queue kuyruğunu, message converter'ı ve listener factory'yi tanımlar.", "Spring AMQP, RabbitMQ"],
            ["CartAuditConsumer", "Consumer", "shopping_cart_queue kuyruğunu dinler ve gelen sepet audit mesajlarını log'a yazar.", "RabbitListener, SLF4J"],
            ["OpenApiConfig", "Configuration", "Swagger/OpenAPI için Bearer JWT güvenlik şemasını ve API Gateway server bilgisini tanımlar.", "Springdoc OpenAPI"],
            ["application.properties", "Configuration", "Servis adı, port, Redis, RabbitMQ, Eureka ve Config Server bağlantılarını tanımlar.", "Spring Boot configuration"],
            ["logback-spring.xml", "Logging", "Konsol ve dosya loglamasını yapılandırır; log dosyasını /tmp veya LOG_DIR altına yazar.", "Logback, SLF4J"],
        ],
    )

    doc.add_heading("4. Kullanılan Teknolojiler ve Projede Kullanım Şekli", level=1)
    add_table(
        doc,
        ["Teknoloji", "Projede nasıl kullanılmış?"],
        [
            ["Java 21", "Servisin geliştirme dili ve runtime sürümü olarak kullanılmış."],
            ["Spring Boot 3.5.14", "REST API, dependency injection, configuration ve uygulama başlatma altyapısı için kullanılmış."],
            ["Spring Web MVC", "Controller endpointleri ve JSON tabanlı HTTP API için kullanılmış."],
            ["Spring Data Redis", "Sepetlerin username anahtarıyla Redis'te saklanması için kullanılmış."],
            ["Redis", "Sepet gibi hızlı değişen ve geçici sayılabilecek veriler için cache/veri deposu olarak kullanılmış."],
            ["RabbitMQ / Spring AMQP", "Sepete ürün ekleme olaylarını shopping_cart_queue kuyruğuna audit mesajı olarak göndermek ve tüketmek için kullanılmış."],
            ["Eureka Client", "Servisin discovery-server'a kaydolması ve diğer servisleri service name ile bulabilmesi için kullanılmış."],
            ["Spring Cloud Config", "Konfigürasyonların merkezi config-server'dan alınabilmesi için kullanılmış."],
            ["RestTemplate + @LoadBalanced", "product-service çağrılarında servis adını kullanarak load-balanced HTTP istekleri atmak için kullanılmış."],
            ["Springdoc OpenAPI / Swagger", "API dokümantasyonu ve JWT bearer güvenlik tanımı için kullanılmış."],
            ["Lombok", "Getter, setter, constructor ve constructor injection kodlarını sadeleştirmek için kullanılmış."],
            ["Logback / SLF4J", "Servis ve RabbitMQ audit mesajlarını loglamak için kullanılmış."],
            ["JUnit 5, Mockito, MockMvc", "Unit, controller ve integration testlerinde kullanılmış."],
            ["Jib Maven Plugin", "Dockerfile yazmadan container image üretimi için yapılandırılmış."],
        ],
    )

    doc.add_heading("5. Veri Modeli", level=1)
    add_table(
        doc,
        ["Model", "Alan", "Açıklama"],
        [
            ["ShoppingCard", "username", "Redis id alanıdır; her kullanıcı için tek sepet mantığı kurar."],
            ["ShoppingCard", "items", "Sepetteki CardItem listesidir."],
            ["ShoppingCard", "totalPrice", "Sepetteki tüm ürünlerin price * quantity toplamıdır."],
            ["CardItem", "productId", "Product Service'teki ürün id'sidir."],
            ["CardItem", "title", "Ürün adı; Product Service'ten alınır."],
            ["CardItem", "quantity", "Sepetteki adet bilgisidir."],
            ["CardItem", "price", "Ürün fiyatıdır; toplam hesaplamada kullanılır."],
            ["CardItem", "imageUrl", "Frontend sepet görünümü için ürün görselidir."],
            ["CardItem", "color", "Ürün rengi; Product Service'te yoksa Standart atanır."],
        ],
    )

    doc.add_heading("6. Testlerde Neler Doğrulanmış?", level=1)
    add_table(
        doc,
        ["Test class", "Doğrulanan davranış"],
        [
            ["ShoppingCardServiceTest", "Mevcut sepetin getirilmesi, başarılı ürün ekleme, ürün bulunamazsa hata, sepet temizleme ve RabbitMQ mesaj gönderimi."],
            ["ShoppingCardControllerTest", "Controller endpointlerinin servis katmanına doğru bağlanması ve temel HTTP cevapları."],
            ["ShoppingCardRepositoryTest", "Redis repository üzerinden sepet kaydetme, bulma ve silme davranışı."],
            ["ShoppingCardIntegrationTest", "MockMvc ile gerçek Spring context içinde sepet getirme, ürün ekleme, sepet temizleme ve aynı ürün tekrar eklenince miktarın artması."],
            ["ShoppingCardServiceApplicationTests", "Spring context'in ayağa kalkabildiğini kontrol eden temel smoke test."],
        ],
    )

    doc.add_heading("7. Mülakatta Anlatım İçin Hazır Kısa Sunum", level=1)
    doc.add_paragraph(
        "Bu serviste kullanıcı sepetini ayrı bir mikroservis olarak tasarladım. Sepet verisi çok sık değiştiği ve siparişe dönüşene kadar geçici bir veri olduğu için Redis kullandım. "
        "Her kullanıcı için username'i key gibi kullanarak tek sepet tutuyorum. Ürün sepete eklendiğinde ürünün fiyat, isim, görsel gibi bilgilerini Product Service'ten RestTemplate ile alıyorum. "
        "Sonra sepette aynı ürün varsa quantity değerini artırıyorum, yoksa yeni CardItem ekliyorum. Her işlemden sonra toplam fiyatı yeniden hesaplayıp sepeti Redis'e kaydediyorum. "
        "Ayrıca RabbitMQ ile shopping_cart_queue kuyruğuna audit mesajı gönderiyorum; CartAuditConsumer bu mesajı tüketip log'a yazıyor. Servis Eureka'ya register oluyor, API Gateway üzerinden route ediliyor ve Swagger/OpenAPI ile dokümante ediliyor."
    )

    doc.add_heading("8. Mülakat Soru-Cevapları", level=1)
    qas = [
        ("1. Shopping Cart Service'in ana sorumluluğu nedir?",
         "Kullanıcının sepet durumunu yönetir. Sepeti getirme, ürün ekleme, miktar güncelleme, ürün silme ve sepeti temizleme işlemlerini yapar. Sipariş oluşturma veya ödeme alma bu servisin sorumluluğu değildir."),
        ("2. Sepet verisini neden PostgreSQL yerine Redis'te tuttun?",
         "Sepet verisi sık değişir ve siparişe dönüşmeden önce geçici bir durumdur. Redis hızlı okuma-yazma sağlar ve kullanıcı bazlı sepet gibi cache/state verileri için uygundur."),
        ("3. Her kullanıcı için sepet nasıl ayrılıyor?",
         "ShoppingCard modelinde username id gibi kullanılıyor. Redis'te her kullanıcı için username üzerinden ayrı bir ShoppingCard tutuluyor."),
        ("4. Ürün sepete eklenirken neden Product Service çağrılıyor?",
         "Sepet servisi ürünün güncel fiyatını, başlığını, görselini ve renk bilgisini kendisi tutmuyor. Bu bilgiler Product Service'ten alınarak sepet satırına ekleniyor."),
        ("5. RestTemplate neden @LoadBalanced tanımlanmış?",
         "Servis URL'inde product-service gibi servis adı kullanıldığında Eureka ve Spring Cloud LoadBalancer üzerinden instance çözümlemesi yapılabilsin diye."),
        ("6. Aynı ürün sepete tekrar eklenirse ne oluyor?",
         "Yeni satır eklenmiyor. Mevcut CardItem bulunuyor ve quantity değeri artırılıyor. Sonra toplam fiyat yeniden hesaplanıyor."),
        ("7. Toplam fiyat nasıl hesaplanıyor?",
         "items listesindeki her ürün için price * quantity hesaplanıyor ve tüm satırlar toplanarak totalPrice set ediliyor."),
        ("8. RabbitMQ bu serviste hangi amaçla kullanılıyor?",
         "Sepete ürün ekleme olayını shopping_cart_queue kuyruğuna audit mesajı olarak göndermek için kullanılıyor. Bu mesaj CartAuditConsumer tarafından tüketilip log'a yazılıyor."),
        ("9. RabbitMQ mesajı iş sürecini devam ettirmek için mi kullanılıyor?",
         "Bu serviste ana iş süreci için kritik değil, daha çok audit/loglama amaçlı. Mesaj gönderimi sepet eklendi olayını görünür hale getiriyor."),
        ("10. CartAuditConsumer neden aynı servisin içinde?",
         "Queue'da mesaj birikmesini önlemek ve audit mesajlarının tüketildiğini göstermek için. Daha büyük sistemde bu consumer ayrı bir audit/log servisine taşınabilir."),
        ("11. Controller ile Service ayrımı neden var?",
         "Controller HTTP request-response yönetir; iş kuralları service içinde kalır. Böylece kod test edilebilir ve katmanlı mimari korunur."),
        ("12. ShoppingCardRepository nasıl çalışıyor?",
         "CrudRepository<ShoppingCard, String> üzerinden Redis'e kaydetme, bulma, silme gibi temel işlemleri sağlar. String id username'dir."),
        ("13. @RedisHash ne işe yarıyor?",
         "ShoppingCard nesnesinin Redis'te hangi hash/namespace altında saklanacağını belirtir. Bu projede Card adı kullanılmış."),
        ("14. DataSourceAutoConfiguration neden exclude edilmiş?",
         "Servis sepet verisini relational database yerine Redis'te tuttuğu için klasik datasource zorunluluğu devre dışı bırakılmış. POM'da JPA dependency bulunduğu için bu exclude uygulamanın gereksiz DB aramasını engelliyor."),
        ("15. API Gateway ile bağlantısı nasıl?",
         "Gateway, /api/shopping-cart/** path'ini SHOPPING-CART-SERVICE'e route ediyor. Böylece frontend doğrudan servise değil gateway'e gider."),
        ("16. Swagger/OpenAPI nasıl kullanılmış?",
         "Controller endpointleri @Operation ve @Tag ile açıklanmış. OpenApiConfig içinde bearerAuth JWT şeması tanımlanmış ve Gateway üzerinden dokümantasyon gösterimi desteklenmiş."),
        ("17. Sepet temizleme ne zaman kullanılır?",
         "Kullanıcı sepeti manuel temizleyebilir veya başarılı sipariş/ödeme akışından sonra order/payment tarafı sepeti temizleyebilir."),
        ("18. Bu serviste güvenlik nasıl düşünülmüş?",
         "OpenAPI'de bearerAuth belirtilmiş ve mimaride API Gateway JWT doğrulaması yapıyor. Ancak endpoint path'inde username almak, token içindeki kullanıcıyla eşleşme kontrolü yapılmazsa güvenlik riski doğurabilir."),
        ("19. Bu serviste iyileştirme olarak ne önerirsin?",
         "Quantity için negatif/sıfır kontrolü eklenebilir, username path yerine JWT içinden alınabilir, Product Service cevabı DTO ile tipli hale getirilebilir, boş ProductRepository silinebilir ve duplicate dependency'ler temizlenebilir."),
        ("20. Product Service kapalıysa ne olur?",
         "addToCart sırasında ürün bilgisi alınamadığı için hata fırlatılır ve sepet kaydedilmez. Daha dayanıklı yapı için timeout, retry, circuit breaker ve anlamlı hata dönüşü eklenebilir."),
        ("21. Redis kapanırsa ne olur?",
         "Sepet okuma-yazma işlemleri başarısız olur. Üretim ortamında Redis persistence, health check ve fallback stratejileri değerlendirilmelidir."),
        ("22. Neden sepeti order-service içinde tutmadın?",
         "Sepet ve sipariş farklı yaşam döngülerine sahip. Sepet geçici ve sık değişen state, sipariş ise ödeme/stok sonrası kalıcı iş kaydıdır. Ayrı servis sorumlulukları daha temiz ayırır."),
        ("23. Unit testlerde neyi mockladın?",
         "Repository, RestTemplate ve RabbitTemplate mocklandı. Böylece service iş kuralları dış sistemlere bağlı kalmadan test edildi."),
        ("24. Integration testlerde neyi doğruladın?",
         "Spring context ve MockMvc ile endpointlerin birlikte çalışması, Redis repository davranışı, ürün ekleme sonrası toplam fiyat ve RabbitMQ gönderim çağrısı doğrulandı."),
        ("25. Bu servis hangi design pattern'lere yakın?",
         "Repository pattern, cache-aside yaklaşımı, event publishing/audit messaging ve katmanlı mimari uygulanmış. Mikroservis mimarisinde Service Discovery ve API Gateway pattern'leriyle çalışıyor."),
    ]
    for q, a in qas:
        add_qa(doc, q, a)

    doc.add_heading("9. Dikkat Edilecek Noktalar ve İyileştirme Fikirleri", level=1)
    add_bullets(
        doc,
        [
            "clearCart controller metodu 204 No Content döndürüyor; bazı testlerde 200 OK ve metin beklenmiş. Test beklentisi ile gerçek controller davranışı hizalanmalı.",
            "quantity için negatif veya sıfır değer kontrolü yok. Sepet toplamını bozabilecek girişler engellenmeli.",
            "username path üzerinden alınıyor. Güvenli mimaride username token claim'inden alınmalı veya path username ile token kullanıcısı karşılaştırılmalı.",
            "Product Service cevabı raw HashMap ile okunuyor. DTO kullanmak tip güvenliği ve okunabilirlik sağlar.",
            "ProductRepository.java boş görünüyor; kullanılmıyorsa kaldırılabilir.",
            "pom.xml içinde bazı dependency'ler tekrar edilmiş; sadeleştirmek bakım kolaylığı sağlar.",
            "RabbitMQ audit mesajı String olarak gidiyor. İleride event DTO + JSON converter ile daha genişletilebilir hale getirilebilir.",
            "Hata yönetiminde genel RuntimeException yerine özel exception ve global exception handler kullanılabilir.",
        ],
    )

    doc.add_heading("10. Kısa Ezber Kartları", level=1)
    add_table(
        doc,
        ["Konu", "Tek cümlelik cevap"],
        [
            ["Redis", "Sepet hızlı değiştiği için username key'i ile Redis'te tutuldu."],
            ["RabbitMQ", "Sepete ürün ekleme olayı audit amacıyla queue'ya publish edildi."],
            ["Product Service", "Ürün bilgisi ve fiyatın güncel kaynağı olduğu için ekleme anında çağrıldı."],
            ["Eureka", "Servislerin birbirini isimle bulması ve gateway routing için kullanıldı."],
            ["API Gateway", "Frontend isteklerini /api/shopping-cart/** üzerinden bu servise yönlendirdi."],
            ["Test", "Unit testlerde mock, integration testlerde MockMvc ile akış doğrulandı."],
        ],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
